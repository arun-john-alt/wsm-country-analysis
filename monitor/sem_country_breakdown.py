#!/usr/bin/env python3
"""
Per-country SEM Brand/Non-Brand breakdown for all NON-presales countries.
Recipe validated against Brazil ground truth (Google NB conv 2024=14/2025=1/H1-2026=1).

Grain: Country x Engine(Google/Bing) x Type(Brand/Non-Brand) x Product x Period(2024/2025/H1-2026)
Metric: conversions = SUM(FS_Valid_Converted_Leads); spend = SUM(Cost) [leads-table Cost is already USD]
Filters: Lead_Type='All Leads', Campaign_Network='Search', Source_Medium in google/bing cpc.
Excludes the 5 presales markets: US, India, UK, Canada, Australia.
"""
import sys
sys.path.insert(0, '/Users/arun-8846/Downloads/Monitor/wsm-monitor')
from wsm_cfg import bq_client
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = '/Users/arun-8846/Downloads/Monitor/SEM_Country_Breakdown.xlsx'
PRESALES = ('United States', 'India', 'United Kingdom', 'Canada', 'Australia')

SQL = """
SELECT
  CampaignCountry AS country,
  CASE WHEN Source_Medium='google / cpc' THEN 'Google' ELSE 'Bing' END AS engine,
  CASE WHEN Theme='Branding' THEN 'Brand' ELSE 'Non-Brand' END AS type,
  Product AS product,
  CASE WHEN SUBSTR(Date,1,4)='2024' THEN 'y2024'
       WHEN SUBSTR(Date,1,4)='2025' THEN 'y2025'
       ELSE 'h1_2026' END AS period,
  SUM(FS_Valid_Converted_Leads) AS conv,
  SUM(Cost) AS spend
FROM `it-security-online-marketing.Google_ads_data_ajay.themes_firstlast_semroi`
WHERE Lead_Type='All Leads'
  AND Campaign_Network='Search'
  AND Source_Medium IN ('google / cpc','bing / cpc')
  AND CampaignCountry IS NOT NULL AND CampaignCountry != '' AND CampaignCountry != '-'
  AND CampaignCountry NOT IN ('United States','India','United Kingdom','Canada','Australia')
  AND (SUBSTR(Date,1,4) IN ('2024','2025') OR SUBSTR(Date,1,7) BETWEEN '2026-01' AND '2026-06')
GROUP BY 1,2,3,4,5
"""

def fetch():
    bq = bq_client()
    # data[(country,engine,type,product)] = {period: {'conv':x,'spend':y}}
    data = {}
    for r in bq.query(SQL).result():
        key = (r['country'], r['engine'], r['type'], r['product'] or '(none)')
        d = data.setdefault(key, {})
        d[r['period']] = {'conv': float(r['conv'] or 0), 'spend': float(r['spend'] or 0)}
    return data

PERIODS = ['y2024', 'y2025', 'h1_2026']
PLABEL = {'y2024': '2024', 'y2025': '2025', 'h1_2026': "H1 2026"}

def cell(d, period, field):
    return d.get(period, {}).get(field, 0)

# ---- styling helpers ----
HDR = Font(bold=True, color='FFFFFF', size=10)
HDRFILL = PatternFill('solid', fgColor='305496')
SUBFONT = Font(bold=True, size=10)
SUBFILL = PatternFill('solid', fgColor='E7EEF7')
GRP = {'Google': 'FFFFFF', 'Bing': 'FBF4E7'}
thin = Side(style='thin', color='D0D0D0')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        x = ws.cell(row=row, column=c)
        x.font = HDR; x.fill = HDRFILL; x.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        x.border = BORDER

def build():
    data = fetch()
    countries = sorted({k[0] for k in data})
    wb = Workbook()

    # ================= DETAIL =================
    ws = wb.active; ws.title = 'Detail'
    cols = ['Country', 'Engine', 'Type', 'Product',
            'Conv 2024', 'Conv 2025', 'Conv H1 2026',
            'Spend 2024', 'Spend 2025', 'Spend H1 2026']
    ws.append(cols); style_header(ws, len(cols))
    r = 2
    for country in countries:
        for engine in ['Google', 'Bing']:
            for typ in ['Non-Brand', 'Brand']:
                grp = {k: v for k, v in data.items() if k[0] == country and k[1] == engine and k[2] == typ}
                if not grp:
                    continue
                # product rows sorted by total conv desc, then spend desc
                def tot(v, f): return sum(cell(v, p, f) for p in PERIODS)
                prods = sorted(grp.items(), key=lambda kv: (-tot(kv[1], 'conv'), -tot(kv[1], 'spend')))
                # subtotal
                sub = {p: {'conv': sum(cell(v, p, 'conv') for _, v in prods),
                           'spend': sum(cell(v, p, 'spend') for _, v in prods)} for p in PERIODS}
                ws.append([country, engine, typ, '▸ ALL',
                           sub['y2024']['conv'], sub['y2025']['conv'], sub['h1_2026']['conv'],
                           sub['y2024']['spend'], sub['y2025']['spend'], sub['h1_2026']['spend']])
                for c in range(1, len(cols) + 1):
                    ws.cell(row=r, column=c).font = SUBFONT
                    ws.cell(row=r, column=c).fill = SUBFILL
                    ws.cell(row=r, column=c).border = BORDER
                r += 1
                for (k, v) in prods:
                    ws.append([country, engine, typ, k[3],
                               cell(v, 'y2024', 'conv'), cell(v, 'y2025', 'conv'), cell(v, 'h1_2026', 'conv'),
                               cell(v, 'y2024', 'spend'), cell(v, 'y2025', 'spend'), cell(v, 'h1_2026', 'spend')])
                    for c in range(1, len(cols) + 1):
                        ws.cell(row=r, column=c).border = BORDER
                    r += 1
    # number formats
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for x in row: x.number_format = '#,##0'
    for row in ws.iter_rows(min_row=2, min_col=8, max_col=10):
        for x in row: x.number_format = '$#,##0'
    widths = [16, 9, 11, 26, 10, 10, 12, 12, 12, 13]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(cols))}1'

    # ================= SUMMARY =================
    sm = wb.create_sheet('Summary')
    scols = ['Country', 'Engine', 'Type', 'Conv 2024', 'Conv 2025', 'Conv H1 2026',
             'Spend 2024', 'Spend 2025', 'Spend H1 2026', 'Cost/Conv 2025', 'Signal']
    sm.append(scols); style_header(sm, len(scols))
    sr = 2
    for country in countries:
        for engine in ['Google', 'Bing']:
            for typ in ['Non-Brand', 'Brand']:
                grp = {k: v for k, v in data.items() if k[0] == country and k[1] == engine and k[2] == typ}
                if not grp:
                    continue
                agg = {p: {'conv': sum(cell(v, p, 'conv') for v in grp.values()),
                           'spend': sum(cell(v, p, 'spend') for v in grp.values())} for p in PERIODS}
                c24, c25, c26 = agg['y2024']['conv'], agg['y2025']['conv'], agg['h1_2026']['conv']
                s24, s25, s26 = agg['y2024']['spend'], agg['y2025']['spend'], agg['h1_2026']['spend']
                cpc25 = (s25 / c25) if c25 else 0
                sig = ''
                if engine == 'Bing':
                    sig = 'Bing conversions not in BQ (spend only)'
                elif c24 >= 5 and (c25 + c26) <= max(2, 0.25 * c24) and s25 >= 20000:
                    sig = 'CONVERSIONS COLLAPSED vs 2024, spend sustained'
                elif c25 == 0 and s25 >= 20000:
                    sig = 'Spend with ZERO conversions (2025)'
                elif c26 == 0 and s26 >= 10000:
                    sig = 'H1-2026 spend, zero conversions'
                sm.append([country, engine, typ, c24, c25, c26, s24, s25, s26, cpc25, sig])
                fill = 'FBF4E7' if engine == 'Bing' else 'FFFFFF'
                for c in range(1, len(scols) + 1):
                    x = sm.cell(row=sr, column=c); x.border = BORDER
                    if engine == 'Bing':
                        x.fill = PatternFill('solid', fgColor=fill)
                if sig.startswith('CONVERSIONS') or sig.startswith('Spend with ZERO'):
                    for c in range(1, len(scols) + 1):
                        sm.cell(row=sr, column=c).fill = PatternFill('solid', fgColor='F8D7DA')
                sr += 1
    for row in sm.iter_rows(min_row=2, min_col=4, max_col=6):
        for x in row: x.number_format = '#,##0'
    for row in sm.iter_rows(min_row=2, min_col=7, max_col=10):
        for x in row: x.number_format = '$#,##0'
    swidths = [16, 9, 11, 10, 10, 12, 12, 12, 13, 14, 42]
    for i, w in enumerate(swidths, 1):
        sm.column_dimensions[get_column_letter(i)].width = w
    sm.freeze_panes = 'A2'
    sm.auto_filter.ref = f'A1:{get_column_letter(len(scols))}1'

    # ================= READ ME =================
    rm = wb.create_sheet('Read me')
    notes = [
        ('WSM SEM Country Breakdown — Brand / Non-Brand', True),
        ('', False),
        ('Scope: all non-presales countries (excludes US, India, UK, Canada, Australia).', False),
        ('Grain: Country x Engine x Type x Product x Period. Periods: full 2024, full 2025, H1 2026 (Jan-Jun).', False),
        ('', False),
        ('Metric definitions (validated exact vs Brazil ground truth: Google Non-Brand conv 2024=14/2025=1/H1-2026=1):', True),
        ('  Conversions = FS_Valid_Converted_Leads  (validated/converted first-source leads; NOT GAds conversions).', False),
        ("  Brand = Theme='Branding';  Non-Brand = every other Theme.", False),
        ("  SEM only: Campaign_Network='Search' (Display excluded).", False),
        ('  Engine split via Source_Medium (google / cpc vs bing / cpc).', False),
        ('  Spend = leads-table Cost, already in USD.', False),
        ('', False),
        ('CAVEATS:', True),
        ('  1. BING CONVERSIONS are NOT in BigQuery (external Microsoft/GAds feed). Bing rows show spend only; conv shown as 0.', False),
        ('  2. Current-year (H1 2026) spend runs ~8-18% light due to the late-data sync lag; prior full years reconcile cleanly.', False),
        ('  3. "Region - *" and "Global" are aggregate buckets, not single countries.', False),
        ('', False),
        ('Product codes: ADAP, ADMP, ADSSP, ELA, LOG360, etc. (ADAP=AP, ADMP=MP, ADSSP=SSP in shorthand).', False),
        ('Next step: RCA pass on the red "CONVERSIONS COLLAPSED" / "Spend with ZERO conversions" rows in Summary.', False),
    ]
    for i, (txt, bold) in enumerate(notes, 1):
        c = rm.cell(row=i, column=1, value=txt)
        if bold: c.font = Font(bold=True)
    rm.column_dimensions['A'].width = 120

    wb.save(OUT)
    return data, countries

def verify(data):
    print('\n=== BRAZIL VERIFICATION (vs ground truth) ===')
    for engine, typ in [('Google', 'Non-Brand'), ('Bing', 'Non-Brand'), ('Google', 'Brand')]:
        grp = {k: v for k, v in data.items() if k[0] == 'Brazil' and k[1] == engine and k[2] == typ}
        for p in PERIODS:
            items = sorted(((k[3], cell(v, p, 'conv')) for k, v in grp.items() if cell(v, p, 'conv') > 0),
                           key=lambda x: -x[1])
            tc = sum(cell(v, p, 'conv') for v in grp.values())
            ts = sum(cell(v, p, 'spend') for v in grp.values())
            pd = ', '.join(f'{k}-{int(c)}' for k, c in items)
            print(f'{engine:6} {typ:9} {PLABEL[p]:8}: conv={int(tc):3}  spend=${ts:>10,.0f}  [{pd}]')

def decliners(data, countries):
    print('\n=== GOOGLE NON-BRAND: worst conversion declines (2024 -> 2025 -> H1 2026) ===')
    rows = []
    for country in countries:
        grp = {k: v for k, v in data.items() if k[0] == country and k[1] == 'Google' and k[2] == 'Non-Brand'}
        if not grp:
            continue
        c24 = sum(cell(v, 'y2024', 'conv') for v in grp.values())
        c25 = sum(cell(v, 'y2025', 'conv') for v in grp.values())
        c26 = sum(cell(v, 'h1_2026', 'conv') for v in grp.values())
        s25 = sum(cell(v, 'y2025', 'spend') for v in grp.values())
        rows.append((country, c24, c25, c26, s25))
    # rank: biggest absolute drop from 2024, still spending in 2025
    rows.sort(key=lambda r: (-(r[1] - r[2]), -r[4]))
    print(f'{"Country":18} {"2024":>5} {"2025":>5} {"H1-26":>6} {"Spend25":>11}')
    for country, c24, c25, c26, s25 in rows[:12]:
        print(f'{country:18} {int(c24):>5} {int(c25):>5} {int(c26):>6} ${s25:>10,.0f}')

if __name__ == '__main__':
    data, countries = build()
    print(f'Countries with SEM activity (non-presales): {len(countries)}')
    print(', '.join(countries))
    verify(data)
    decliners(data, countries)
    print(f'\nWorkbook: {OUT}')
